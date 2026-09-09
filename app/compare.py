"""Deterministic comparison and formula-safe Excel export."""
import io,re
from decimal import Decimal
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from .extract import FIELDS,STAGES,METRICS,clean,numeric
from .quantities import quantity_key
from datetime import datetime

COLORS={'same':'FFFFFF','different':'FFF0D7','equivalent':'E7F1FF','missing':'EEF1F5','review':'FCE7E8'}
LABELS={'same':'一致','different':'内容变化','equivalent':'表述不同 / 标准值相同','missing':'未载明','review':'待核对'}

def norm(value):return clean(value).replace('〔','[').replace('〕',']').replace('，',',').replace('。','')
def compare_cells(cells):
    values=[c.get('value') for c in cells]
    present=[v for v in values if v is not None and v!='']
    if any(c.get('status') in ['needs_review','conflict','uncertain'] for c in cells):status='review'
    elif len(present)!=len(values) and len(set(present))<=1:status='missing'
    elif len(set(values))<=1:status='same'
    elif len({norm(v) for v in present})==1:status='equivalent'
    else:status='different'
    return status

def date_key(value):
    m=re.search(r'(\d{4})年(\d{1,2})月(\d{1,2})日',value or '')
    return tuple(map(int,m.groups())) if m else (0,0,0)

def rows_for(docs):
    docs=sorted(docs,key=lambda d:(STAGES.index(d['stage']) if d['stage'] in STAGES else 3,date_key(d['fields']['印发日期'].get('value')),d['filename']))
    rows=[]
    for name in FIELDS:
        cs=[d['fields'][name] for d in docs]
        status=compare_cells(cs);note=''
        if name in ['建设周期','总投资/匡算/估算/概算'] and status not in ['review','missing']:
            ns=[numeric(c.get('value')) for c in cs]
            bases=[re.search(r'匡算|估算|概算',c.get('value') or '') for c in cs]
            basis=[m.group() if m else '总投资' for m in bases]
            if all(ns) and len({quantity_key(n) for n in ns})==1:
                if name=='建设周期' or len(set(basis))==1:
                    status='equivalent' if len({c['value'] for c in cs})>1 else status
            if name=='总投资/匡算/估算/概算' and len(set(basis))>1:note='投资口径：'+' → '.join(basis)+'；阶段金额变化不直接认定为异常。'
        rows.append({'name':name,'kind':'fixed','cells':cs,'status':status,'note':note})
    names=list(dict.fromkeys(m['name'] for d in docs for m in d['metrics']))
    priority={name:i for i,(name,_) in enumerate(METRICS)}
    names.sort(key=lambda name:priority.get(name,len(priority)))
    for name in names:
        cs=[]
        for d in docs:
            hits=[m for m in d['metrics'] if m['name']==name]
            if not hits:cs.append({'value':None,'status':'missing','evidence':[]})
            elif len(hits)==1:cs.append(hits[0])
            else:cs.append({'value':'；'.join(m['value'] for m in hits),'status':'conflict','evidence':[e for m in hits for e in m['evidence']]})
        status=compare_cells(cs);note=''
        ns=[numeric(c['value']) if c.get('value') else None for c in cs]
        available=[(i,n) for i,n in enumerate(ns) if n is not None]
        if len(available)>1 and status!='review':
            original_ns=ns
            indices=[i for i,n in available]
            ns=[n for i,n in available]
            keys=[quantity_key(n) for n in ns]
            if len(set(keys))==1 and len({c['value'] for c in cs})>1:status='equivalent' if len(available)==len(cs) else 'missing'
            if len({n['unit'] for n in ns})==1 and not any(n['upper'] for n in ns):
                notes=[]
                for i in range(1,len(ns)):
                    a,b=Decimal(ns[i-1]['number']),Decimal(ns[i]['number'])
                    if a!=b:
                        note_i=f'{docs[indices[i-1]]["stage"]}→{docs[indices[i]]["stage"]}：{b-a:+f}{ns[i]["unit"]}'
                        if a:
                            pct=(b-a)/a*100
                            note_i+=('（变动比例小于0.01%）' if abs(pct)<Decimal('.005') else f'（{pct:+.2f}%）')
                        if ns[i-1]['qualifier'] or ns[i]['qualifier']:note_i+='，含约数/限定'
                        notes.append(note_i)
                note='；'.join(notes)
        rows.append({'name':name,'kind':'metric','cells':cs,'status':status,'note':note})
    return docs,rows

def safe(value):
    s='' if value is None else str(value)
    return "'"+s if s.lstrip().startswith(('=','+','-','@')) else s

def export_xlsx(documents):
    wb=Workbook();wb.remove(wb.active)
    groups={}
    for d in documents:groups.setdefault(d['project_key'],[]).append(d)
    evidence=wb.create_sheet('证据索引');evidence.append(['项目','文件','阶段','字段','值','PDF页码','原文','坐标'])
    single=wb.create_sheet('单文档结构化');single.append(['项目','文件','阶段','字段','提取值','状态','证据页码'])
    history=wb.create_sheet('修订记录');history.append(['文件','时间','字段','修改前','修改后','说明'])
    review=wb.create_sheet('待复核事项');review.append(['项目','文件','事项'])
    for gi,ds in enumerate(groups.values(),1):
        docs,rows=rows_for(ds)
        for kind,label in [('fixed','固定字段'),('metric','建设指标')]:
            ws=wb.create_sheet(f'{gi}-{label}')
            ws.append([safe(docs[0]['project_name'])]);ws.append(['字段']+[d['stage']+' | '+d['filename'] for d in docs]+['比较结果','说明'])
            for r in rows:
                if r['kind']!=kind:continue
                ws.append([safe(r['name'])]+[safe(c.get('value') if c.get('value') is not None else '未载明') for c in r['cells']]+[LABELS[r['status']],safe(r['note'])])
                for c in ws[ws.max_row]:c.fill=PatternFill('solid',fgColor=COLORS['missing' if c.value=='未载明' else r['status']])
            ws.freeze_panes='B3';ws.auto_filter.ref=f'A2:{ws.cell(ws.max_row,ws.max_column).coordinate}'
        for d in docs:
            for k,c in list(d['fields'].items())+[(m['name'],m) for m in d['metrics']]:
                single.append([safe(d['project_name']),safe(d['filename']),d['stage'],safe(k),safe(c.get('value')),c.get('status'),','.join(str(p) for p in sorted({e['page'] for e in c['evidence']}))])
                for e in c['evidence']:
                    evidence.append([safe(d['project_name']),safe(d['filename']),d['stage'],safe(k),safe(c.get('value')),e['page'],safe(e['quote']),str(e['bbox'])])
                if c.get('status') in ['needs_review','conflict','uncertain']:
                    review.append([safe(d['project_name']),safe(d['filename']),safe(k+'：待核对')])
            for h in d.get('history',[]):
                history.append([safe(d['filename']),datetime.fromtimestamp(h['time']).isoformat(),safe(h.get('name','')),safe(h.get('before',{}).get('value') if isinstance(h.get('before'),dict) else h.get('before')),safe(h.get('after')),safe(h.get('reason',''))])
            for w in d['warnings']:review.append([safe(d['project_name']),safe(d['filename']),safe(w)])
    for ws in wb:
        for row in ws:
            for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
        head=2 if '-' in ws.title else 1
        for c in ws[head]:c.fill=PatternFill('solid',fgColor='173A60');c.font=Font(color='FFFFFF',bold=True)
        for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=38 if col[0].column>1 else 25
        if ws.title=='证据索引':ws.column_dimensions['G'].width=75
    out=io.BytesIO();wb.save(out);return out.getvalue()
